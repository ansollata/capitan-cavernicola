#!/usr/bin/env python3
"""Generate the proposed Premiere Networks (PN) rate card (v0.1 DRAFT) from
the ORIGINAL `PN Rate Cards.xlsx` (Updated Rates 6.1 folder).

Design (docs/05-ratecard-solution.md section 2.3): CONSOLIDATION, not
re-derivation. The source spreads the same Vehicle x Daypart x length across
22 overlapping price books. The proposal: ONE base card + named discount
schedules expressed as % of base, with dead rows dropped and the flat
quarterly columns collapsed to a single rate.

Usage: python3 generate_pn_card.py <PN.xlsx> <output.xlsx>
"""

import statistics as st
import sys
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------- KNOBS (management decisions) -----------------
# K1. Book priority for sourcing the BASE rate of each Vehicle/Daypart/len.
#     General-market base first, then ethnic-family bases, then voiced/nets.
BASE_PRIORITY = [
    "2026 Rate Card", "2026 AA Rate Card", "2026 HA Rate Card",
    "2026 Voiced Rate Card", "2026 Voiced AA Rate Card",
    "2026 Voiced HA Rate Card", "2026 Voiced Talk Nets",
    "AA Jelli Nets", "HA Jelli Nets",
]
# Books that become DISCOUNT SCHEDULES (sales variants of the same
# inventory), never base-rate sources.
SCHEDULE_BOOKS_HINT = ("TARGET", "Label")

# K2. Quarter collapse: median of non-zero quarters (data shows quarters are
#     flat for most rows); rows whose quarters vary beyond this are flagged.
QTR_VARY_FLAG = 0.10

# K3. Round pricing.
def round_price(x):
    if x <= 0:
        return 0
    if x < 20:
        return max(1, round(x))
    if x < 100:
        return int(5 * round(x / 5))
    return int(10 * round(x / 10))

# K4. Discount-schedule consistency: a book whose row-level ratios to base
#     have IQR wider than this is flagged "not-a-clean-schedule".
SCHEDULE_IQR_FLAG = 0.15
# -----------------------------------------------------------------------------


def main(src, out_path):
    ws = load_workbook(src, read_only=True)["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = {h: k for k, h in enumerate(rows[0])}
    qs = ["Qtr1", "Qtr2", "Qtr3", "Qtr4"]

    live = []
    dead = 0
    for r in rows[1:]:
        qvals = [r[hdr[q]] for q in qs]
        qvals = [v for v in qvals if isinstance(v, (int, float)) and v > 0]
        if not qvals:
            dead += 1
            continue
        rate = st.median(qvals)
        varies = (max(qvals) - min(qvals)) / rate > QTR_VARY_FLAG if rate else False
        live.append({
            "book": r[hdr["RateDesc"]], "veh": r[hdr["Vehicle"]],
            "dp": r[hdr["Daypart"]], "len": r[hdr["len"]],
            "rate": rate, "qvary": varies,
        })

    by_key = defaultdict(dict)  # (veh, dp, len) -> book -> row
    for x in live:
        by_key[(x["veh"], x["dp"], x["len"])][x["book"]] = x

    prio = {b: k for k, b in enumerate(BASE_PRIORITY)}
    base = {}
    for key, books in by_key.items():
        candidates = [b for b in books if b in prio]
        if candidates:
            src_book = min(candidates, key=lambda b: prio[b])
            base[key] = (books[src_book], src_book)

    # Discount schedules: every non-base-source book vs the base rate.
    sched = defaultdict(list)
    for key, books in by_key.items():
        if key not in base:
            continue
        brate = base[key][0]["rate"]
        if brate <= 0:
            continue
        for book, x in books.items():
            if book != base[key][1]:
                sched[book].append(x["rate"] / brate)

    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="C6002B")
    flag_fill = PatternFill("solid", fgColor="FFF2CC")

    w = wb.active
    w.title = "Proposed PN Card v0.1 DRAFT"
    cols = ["Vehicle", "Daypart", "Len", "Source Book", "Cur Rate", "PROP Rate", "Flags"]
    w.append(cols)
    for c in w[1]:
        c.font, c.fill = hdr_font, hdr_fill
    n_flagged = 0
    for key in sorted(base, key=lambda k: (str(k[0]), str(k[1]), k[2])):
        x, src_book = base[key]
        flags = []
        if src_book != "2026 Rate Card":
            flags.append(f"base-from-variant-book")
        if x["qvary"]:
            flags.append("quarters-vary>10%")
        prop = round_price(x["rate"])
        w.append([key[0], key[1], key[2], src_book, round(x["rate"], 2), prop,
                  ";".join(flags)])
        if flags:
            w.cell(row=w.max_row, column=len(cols)).fill = flag_fill
            n_flagged += 1
    w.freeze_panes = "D2"
    w.auto_filter.ref = w.dimensions
    for idx, width in enumerate([42, 16, 5, 24, 10, 10, 34], 1):
        w.column_dimensions[get_column_letter(idx)].width = width

    d = wb.create_sheet("Discount Schedules")
    d.append(["Book (RateDesc)", "Rows vs base", "Median % of base",
              "p25 %", "p75 %", "Assessment"])
    for c in d[1]:
        c.font, c.fill = hdr_font, hdr_fill
    for book in sorted(sched, key=lambda b: -len(sched[b])):
        v = sorted(sched[book])
        n = len(v)
        med, p25, p75 = st.median(v), v[n // 4], v[3 * n // 4]
        clean = (p75 - p25) <= SCHEDULE_IQR_FLAG
        d.append([book, n, round(100 * med, 1), round(100 * p25, 1),
                  round(100 * p75, 1),
                  "clean schedule" if clean else "NOT clean - review row variance"])
        if not clean:
            d.cell(row=d.max_row, column=6).fill = flag_fill
    for idx, width in enumerate([34, 12, 16, 8, 8, 30], 1):
        d.column_dimensions[get_column_letter(idx)].width = width

    s = wb.create_sheet("Summary")
    for line in [
        ("Source rows", len(rows) - 1),
        ("Dead all-zero rows dropped", dead),
        ("Live rows", len(live)),
        ("Cells on proposed base card (Vehicle x Daypart x Len)", len(base)),
        ("  of which flagged", n_flagged),
        ("Books in source", len({x['book'] for x in live}) ),
        ("Books converted to discount schedules", len(sched)),
        ("Live rows NOT on base card (no base-priority book for key)",
         sum(len(b) for k, b in by_key.items() if k not in base)),
    ]:
        s.append(list(line))
    s.column_dimensions["A"].width = 55
    s["A1"].font = Font(bold=True)

    m = wb.create_sheet("Methodology & Knobs")
    for t in [
        "PROPOSED PREMIERE (PN) RATE CARD v0.1 - DRAFT FOR MANAGEMENT REVIEW",
        "",
        "SOURCE: original 'PN Rate Cards.xlsx' (Updated Rates 6.1), firsthand.",
        "",
        "DESIGN: consolidation. 22 overlapping price books collapse to ONE base",
        "card (Vehicle x Daypart x Length, single rate) plus NAMED DISCOUNT",
        "SCHEDULES expressed as % of base. Dead all-zero rows dropped. Quarterly",
        "columns collapsed to one rate (median of non-zero quarters) because the",
        "data shows no quarterly structure; rows whose quarters do vary are",
        "flagged 'quarters-vary>10%' for review.",
        "",
        "KNOBS:",
        f"K1 base-book priority: {BASE_PRIORITY}",
        f"K2 quarter-variation flag threshold: {int(QTR_VARY_FLAG*100)}%",
        "K3 rounding: $1 under 20, $5 under 100, $10 above",
        f"K4 schedule cleanliness: IQR of row ratios <= {SCHEDULE_IQR_FLAG} of base",
        "",
        "MANAGEMENT DECISIONS THIS SETS UP:",
        "- Approve one base card; retire per-book rate maintenance.",
        "- For each discount schedule: keep (as a named % off base), merge, or kill.",
        "- Schedules marked 'NOT clean' are books whose discounts are inconsistent",
        "  row to row - decide the intended % and normalize.",
    ]:
        m.append([t])
    m.column_dimensions["A"].width = 95

    wb.save(out_path)
    print(f"live={len(live)} dead={dead} base_cells={len(base)} schedules={len(sched)}")
    print("saved:", out_path)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
